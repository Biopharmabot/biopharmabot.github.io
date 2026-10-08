# -*- coding: utf-8 -*-
"""13F 헷지펀드 분기 산출물(D:\\13F\\out\\Hedge_fund_<분기>.xlsx) → data/hedge.json
사용: python tools/export_hedge.py [2Q26]   (생략 시 out/ 최신 파일)
페이지의 '헷지펀드' 탭·개요 카드·회사 모아보기가 읽는다. 분기마다 13F 파이프라인(run_all.py) 뒤에 한 번 실행하고 커밋.
"""
import sys, json, re, glob, os, datetime as dt
sys.stdout.reconfigure(encoding="utf-8")
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
F13 = r"D:\13F"
FUNDS = ["Baker Bros", "RA Capital", "Orbimed", "RTW", "Avoro", "Deerfield", "Perceptive", "Deep Track"]


def qkey(q):
    return (int("20" + q[2:]), int(q[0]))


def load_fund(ws):
    """펀드 시트 7행~: Rank·Stock·Issuer·Class·CUSIP·Shares·MV($M)… → 티커별 합산(복수 클래스 합산)"""
    rows = {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        t, issuer, cls, sh, mv = r[1], r[2], r[3], r[5], r[6]
        if not t or mv is None:
            continue
        t = str(t).strip()
        if t.startswith("[") or (cls and re.search(r"NOTE|BOND|CALL|PUT|WARRANT", str(cls), re.I)):
            continue  # 채권·옵션·티커 미상은 분석 제외(파이프라인 규칙)
        try:
            sh_n, mv_n = int(float(sh or 0)), float(mv or 0)
        except (TypeError, ValueError):
            continue  # '(bond $125,000,000)' 같은 채권 표기
        d = rows.setdefault(t, {"t": t, "issuer": str(issuer or "").strip(), "sh": 0, "mv": 0.0})
        d["sh"] += sh_n
        d["mv"] += mv_n
    tot = sum(d["mv"] for d in rows.values()) or 1
    for d in rows.values():
        d["wt"] = d["mv"] / tot
    return rows, tot


KEY = [("신규 편입", "new"), ("순매수", "add"), ("순매도", "trim"), ("혼조", "mixed"), ("완전 청산", "exit"),
       ("M&A 소멸", "ma"), ("순확대", "add"), ("순축소", "trim"), ("주석", "note"), ("참고", "skip")]


def parse_groups(ws):
    """동향 요약 시트의 목록 블록을 {기준:{구분:[{n, items}]}} 로"""
    out = {"shares": {}, "weight": {}}
    basis, sect = None, None
    for r in ws.iter_rows(min_row=9, values_only=True):
        a = str(r[0]).strip() if r[0] is not None else ""
        b = str(r[1]).strip() if r[1] is not None else ""
        if not a:
            continue
        if a.startswith("■ 주식"):
            basis = "shares"; continue
        if a.startswith("■ 금액"):
            basis = "weight"; continue
        if not basis:
            continue
        hit = next((v for k, v in KEY if a.startswith(k)), None)
        if hit and not b:
            sect = hit; continue
        if a.startswith("주석"):
            out[basis].setdefault("note", []).append(b); continue
        if a.startswith("참고") or sect in (None, "skip", "note") or not b:
            continue
        m = re.match(r"(\d+)개 펀드", a)
        n = int(m.group(1)) if m else None
        text = re.sub(r"\s+←.*$", "", b)  # 설명 꼬리 제거
        items = [x.strip() for x in text.split(",") if x.strip()] if sect != "ma" else [text]
        out[basis].setdefault(sect, []).append({"n": n, "items": items})
    return out


def bare(x):
    return re.sub(r"\(.*?\)|\*", "", x).strip()


def main(argv):
    files = sorted(glob.glob(os.path.join(F13, "out", "Hedge_fund_*.xlsx")),
                   key=lambda p: qkey(re.search(r"_(\dQ\d\d)", p).group(1)))
    q = argv[1] if len(argv) > 1 else re.search(r"_(\dQ\d\d)", files[-1]).group(1)
    path = os.path.join(F13, "out", f"Hedge_fund_{q}.xlsx")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    quarters = sorted({re.search(r"(\dQ\d\d)$", w.title).group(1) for w in wb.worksheets
                       if re.search(r"(\dQ\d\d)$", w.title)}, key=qkey)
    prev = quarters[quarters.index(q) - 1] if quarters.index(q) > 0 else None

    funds, matrix = [], {}
    for name in FUNDS:
        cur, tot = load_fund(wb[f"{name} {q}"])
        old, _ = load_fund(wb[f"{name} {prev}"]) if prev else ({}, 0)
        top = sorted(cur.values(), key=lambda d: -d["mv"])
        rows = []
        for d in top[:15]:
            o = old.get(d["t"])
            qoq = "new" if not o else (round(d["sh"] / o["sh"] - 1, 3) if o["sh"] else None)
            rows.append({"t": d["t"], "issuer": d["issuer"], "mv": round(d["mv"], 1), "wt": round(d["wt"], 4), "qoq_sh": qoq})
        funds.append({"name": name, "mv": round(tot, 1), "n": len(cur), "top": rows})
        for d in cur.values():
            m = matrix.setdefault(d["t"], {"t": d["t"], "issuer": d["issuer"], "w": {}, "wp": {}})
            m["w"][name] = round(d["wt"], 4)
        for d in old.values():
            m = matrix.setdefault(d["t"], {"t": d["t"], "issuer": d["issuer"], "w": {}, "wp": {}})
            m["wp"][name] = round(d["wt"], 4)
    mat = []
    for m in matrix.values():
        if not m["w"]:
            continue
        s, sp = sum(m["w"].values()), sum(m["wp"].values())
        mat.append({"t": m["t"], "issuer": m["issuer"], "w": m["w"], "n": len(m["w"]),
                    "sum": round(s, 4), "sum_prev": round(sp, 4), "d": round(s - sp, 4)})
    mat.sort(key=lambda x: -x["sum"])

    sig = []
    for r in wb["Signal — Sector"].iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        sig.append({"t": str(r[0]), "issuer": str(r[1] or ""), "held": r[2] or 0, "add": r[3] or 0, "new": r[4] or 0,
                    "trim": r[5] or 0, "exit": r[6] or 0, "net": r[7] or 0, "wtd": round(float(r[8] or 0), 4)})
    groups = parse_groups(wb["동향 요약"])
    # 매집 = 주식수 순매수 ≥2 ∪ 비중 순확대 ≥3 ∪ 신규 ≥3 (D:\13F CLAUDE.md Hub 연계 규칙)
    accum, dist = set(), set()
    for g in groups["shares"].get("add", []):
        if g["n"] and g["n"] >= 2:
            accum.update(bare(x) for x in g["items"])
    for g in groups["shares"].get("new", []):
        if g["n"] and g["n"] >= 3:
            accum.update(bare(x) for x in g["items"])
    for g in groups["weight"].get("add", []):
        if g["n"] and g["n"] >= 3:
            accum.update(bare(x) for x in g["items"])
    for g in groups["shares"].get("trim", []) + groups["weight"].get("trim", []):
        if g["n"] and g["n"] >= 3:
            dist.update(bare(x) for x in g["items"])
    for g in groups["shares"].get("exit", []):
        if g["n"] and g["n"] >= 2:
            dist.update(bare(x) for x in g["items"])

    out = {"quarter": q, "prev": prev, "quarters": quarters, "generated": dt.date.today().isoformat(),
           "source": "13f.info 분기 보고(분기말+45일 공시) · D:\\13F 파이프라인", "funds": funds,
           "signals": sig, "matrix": mat[:120], "n_universe": len(mat), "n_multi": sum(1 for m in mat if m["n"] >= 2), "groups": groups,
           "accum": sorted(accum), "distrib": sorted(dist)}
    p = os.path.join(ROOT, "data", "hedge.json")
    os.makedirs(os.path.dirname(p), exist_ok=True)
    json.dump(out, open(p, "w", encoding="utf-8"), ensure_ascii=False, separators=(",", ":"))
    print(f"{q} (prev {prev}) · funds {len(funds)} · signals {len(sig)} · matrix {len(mat)} · "
          f"accum {len(accum)} · distrib {len(dist)} → {p} ({os.path.getsize(p)//1024} KB)")


if __name__ == "__main__":
    main(sys.argv)
